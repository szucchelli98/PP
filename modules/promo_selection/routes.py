from io import BytesIO

from flask import Blueprint, Response, jsonify, render_template, request

promo_bp = Blueprint("promo_selection", __name__, template_folder="templates")

# ── Discount column detection ─────────────────────────────────────────────────
DISCOUNT_SUFFIX = " discount"

STOCK_COLUMNS = [
    "TOTAL STOCK (ALL REGIONS)",
    "EU STOCK",
    "US STOCK",
    "RU STOCK",
    "ASIA STOCK",
    "AMERICAS STOCK",
    "ITG STOCK",
    "SPEDIMEX STOCK",
    "EU DOS STOCK",
    "DISPLAY ZONE STOCK",
]




# ── Routes ────────────────────────────────────────────────────────────────────

@promo_bp.route("/")
def index():
    return render_template("promo_selection/index.html")


@promo_bp.route("/parse", methods=["POST"])
def parse():
    """Stream the xlsx with openpyxl read-only to avoid OOM on free tier."""
    file = request.files.get("file")
    if not file:
        return jsonify({"error": "No file provided"}), 400
    try:
        import openpyxl

        wb = openpyxl.load_workbook(BytesIO(file.read()), read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.rows

        # ── Header row ────────────────────────────────────────────────────────
        headers = [cell.value for cell in next(rows_iter)]

        cat_idx   = headers.index("CATEGORY")      if "CATEGORY"      in headers else None
        micro_idx = headers.index("MICROCATEGORY") if "MICROCATEGORY" in headers else None

        # Channels: headers ending with " discount"
        channels = [h[:-len(DISCOUNT_SUFFIX)] for h in headers
                    if isinstance(h, str) and h.lower().endswith(DISCOUNT_SUFFIX)]

        # Stock columns: known names present in headers
        stock_cols = [c for c in STOCK_COLUMNS if c in headers]

        # ── Stream data rows ──────────────────────────────────────────────────
        categories_set      = set()
        microcategories_set = set()

        for row in rows_iter:
            if cat_idx is not None:
                v = row[cat_idx].value
                if v:
                    categories_set.add(str(v))
            if micro_idx is not None:
                v = row[micro_idx].value
                if v:
                    microcategories_set.add(str(v))

        wb.close()

        return jsonify({
            "channels":        channels,
            "categories":      sorted(categories_set),
            "microcategories": sorted(microcategories_set),
            "stock_columns":   stock_cols,
        })
    except Exception as exc:
        return jsonify({"error": str(exc)}), 400


@promo_bp.route("/export", methods=["POST"])
def export():
    file = request.files.get("file")
    if not file:
        return jsonify({"error": "No file provided"}), 400

    channel          = request.form.get("channel", "").strip()
    adjustment_type  = request.form.get("adjustment_type", "none")
    adjustment_value = _float(request.form.get("adjustment_value"))
    min_final_price  = _float(request.form.get("min_final_price"))
    max_final_price  = _float(request.form.get("max_final_price"))
    min_channel_disc = _float(request.form.get("min_channel_discount"))
    max_channel_disc = _float(request.form.get("max_channel_discount"))
    stock_col        = request.form.get("stock_column", "").strip()
    exclude_zero_stock = request.form.get("exclude_zero_stock") == "true"
    include_cats     = set(_list(request.form.get("include_categories")))
    exclude_cats     = set(_list(request.form.get("exclude_categories")))
    include_micros   = set(_list(request.form.get("include_microcategories")))
    exclude_micros   = set(_list(request.form.get("exclude_microcategories")))

    try:
        import openpyxl
        from openpyxl import Workbook

        disc_col_name = f"{channel}{DISCOUNT_SUFFIX}"

        # ── Stream input ──────────────────────────────────────────────────────
        wb_in = openpyxl.load_workbook(BytesIO(file.read()), read_only=True, data_only=True)
        ws_in = wb_in.active
        rows_iter = ws_in.rows

        headers = [cell.value for cell in next(rows_iter)]

        if disc_col_name not in headers:
            wb_in.close()
            return jsonify({"error": f"Discount column '{disc_col_name}' not found."}), 400

        # Column indices
        idx = {h: i for i, h in enumerate(headers) if h is not None}
        disc_idx  = idx.get(disc_col_name)
        price_idx = idx.get("PRICE")
        cat_idx   = idx.get("CATEGORY")
        micro_idx = idx.get("MICROCATEGORY")
        stock_idx = idx.get(stock_col) if stock_col else None

        # Output headers: insert FINAL PRICE after PRICE
        out_headers = list(headers)
        final_price_pos = (price_idx + 1) if price_idx is not None else len(out_headers)
        out_headers.insert(final_price_pos, "FINAL PRICE")

        # ── Build output workbook ─────────────────────────────────────────────
        wb_out = Workbook(write_only=True)
        ws_out = wb_out.create_sheet()
        ws_out.append(out_headers)

        row_count = 0

        for row in rows_iter:
            vals = [cell.value for cell in row]
            # Pad short rows
            while len(vals) < len(headers):
                vals.append(None)

            # ── Compute final price ───────────────────────────────────────────
            price = _to_float(vals[price_idx]) if price_idx is not None else 0.0
            disc  = _to_float(vals[disc_idx])  if disc_idx  is not None else 0.0

            if adjustment_type == "new_discount" and adjustment_value is not None:
                final = price * (1 - adjustment_value / 100.0)
            elif adjustment_type == "extra_discount" and adjustment_value is not None:
                final = price * (1 - disc / 100.0) * (1 - adjustment_value / 100.0)
            else:
                final = price * (1 - disc / 100.0)
            final = round(final, 2)

            # ── Apply filters ─────────────────────────────────────────────────
            if min_final_price  is not None and final < min_final_price:  continue
            if max_final_price  is not None and final > max_final_price:  continue
            if min_channel_disc is not None and disc  < min_channel_disc: continue
            if max_channel_disc is not None and disc  > max_channel_disc: continue

            if exclude_zero_stock and stock_idx is not None:
                if _to_float(vals[stock_idx]) <= 0: continue

            cat = str(vals[cat_idx]) if cat_idx is not None and vals[cat_idx] else ""
            if include_cats  and cat not in include_cats:  continue
            if exclude_cats  and cat in exclude_cats:      continue

            micro = str(vals[micro_idx]) if micro_idx is not None and vals[micro_idx] else ""
            if include_micros and micro not in include_micros: continue
            if exclude_micros and micro in exclude_micros:     continue

            # ── Write row ─────────────────────────────────────────────────────
            out_row = list(vals)
            out_row.insert(final_price_pos, final)
            ws_out.append(out_row)
            row_count += 1

        wb_in.close()

        buf = BytesIO()
        wb_out.save(buf)
        xlsx_bytes = buf.getvalue()

        base = file.filename.rsplit(".", 1)[0]
        return Response(
            xlsx_bytes,
            status=200,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f'attachment; filename="{base}_promo_selection.xlsx"',
                "Content-Length": str(len(xlsx_bytes)),
                "X-Row-Count": str(row_count),
            },
        )
    except Exception as exc:
        return jsonify({"error": str(exc)}), 400


# ── Helpers ───────────────────────────────────────────────────────────────────

def _to_float(val) -> float:
    """Convert a cell value to float, defaulting to 0."""
    try:
        return float(val) if val is not None else 0.0
    except (TypeError, ValueError):
        return 0.0


def _float(val) -> float | None:
    try:
        return float(val) if val not in (None, "", "null") else None
    except (TypeError, ValueError):
        return None


def _list(val) -> list[str]:
    """Parse a JSON-encoded list or comma-separated string into a Python list."""
    if not val:
        return []
    val = val.strip()
    if val.startswith("["):
        import json
        try:
            return [str(v) for v in json.loads(val)]
        except Exception:
            pass
    return [v.strip() for v in val.split(",") if v.strip()]
